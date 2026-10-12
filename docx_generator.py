import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from models import Student, CumulativeReport

EMERALD = RGBColor(27, 67, 50)       # #1b4332
GOLD = RGBColor(180, 142, 60)        # #b48e3c
DARK_GRAY = RGBColor(55, 65, 81)     # #374151
LIGHT_BG = "F3F4F6"
GREEN_BG = "D1FAE5"
AMBER_BG = "FEF3C7"
RED_BG = "FEE2E2"

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_student_docx(student: Student) -> io.BytesIO:
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header / Title Block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = title_p.add_run("STUDENT COUNSELLING & GUIDANCE DEPARTMENT\n")
    run_inst.font.name = "Georgia"
    run_inst.font.size = Pt(16)
    run_inst.font.bold = True
    run_inst.font.color.rgb = EMERALD

    run_sub = title_p.add_run("ADOLESCENT DEVELOPMENT & EVALUATION REPORT\n")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = GOLD

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(8)
    p_div_run = p_div.add_run("―" * 58)
    p_div_run.font.color.rgb = GOLD
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Student Demographics Card
    display_name = student.name if not student.is_confidential else f"Candidate {student.id} (CONFIDENTIAL)"
    display_place = student.place if not student.is_confidential else "[Confidential / Masked]"
    display_parents = f"Father: {student.fathers_occupation} | Mother: {student.mothers_occupation}" if not student.is_confidential else "[Family Profile Confidential]"

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    t_data = [
        ("Student Name:", display_name, "Age:", f"{student.age} Years"),
        ("Origin / Hometown:", display_place, "Program Batch:", student.batch),
        ("Family Context:", display_parents, "Triage Status:", student.report.triage_status if student.report else "Evaluated")
    ]

    for row_idx, (l1, v1, l2, v2) in enumerate(t_data):
        row = table.rows[row_idx]
        cell_a, cell_b = row.cells[0], row.cells[1]
        
        set_cell_background(cell_a, "F9FAFB")
        set_cell_background(cell_b, "F9FAFB")
        set_cell_margins(cell_a, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_b, top=60, bottom=60, left=100, right=100)

        p_a = cell_a.paragraphs[0]
        r_l1 = p_a.add_run(f"{l1} ")
        r_l1.bold = True
        r_l1.font.size = Pt(9.5)
        r_v1 = p_a.add_run(v1)
        r_v1.font.size = Pt(9.5)

        p_b = cell_b.paragraphs[0]
        r_l2 = p_b.add_run(f"{l2} ")
        r_l2.bold = True
        r_l2.font.size = Pt(9.5)
        r_v2 = p_b.add_run(v2)
        r_v2.font.size = Pt(9.5)
        if "Requires Follow-up" in v2 or "🔴" in v2:
            r_v2.font.color.rgb = RGBColor(185, 28, 28)
            r_v2.bold = True
        elif "Well-Adapted" in v2 or "🟢" in v2:
            r_v2.font.color.rgb = RGBColor(21, 128, 61)
            r_v2.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Function to add styled section
    def add_section(heading: str, content: str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
        run_h = h.add_run(heading)
        run_h.font.name = "Georgia"
        run_h.font.size = Pt(11.5)
        run_h.font.bold = True
        run_h.font.color.rgb = EMERALD

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        run_c = p.add_run(content)
        run_c.font.name = "Calibri"
        run_c.font.size = Pt(10.5)
        run_c.font.color.rgb = DARK_GRAY

    if student.report:
        add_section("1. Executive Summary & Program Adaptation", student.report.executive_summary)
        add_section("2. Emotional Profile & Self-Concept", student.report.emotional_profile)
        add_section("3. Social & Family Ecosystem", student.report.social_family_dynamics)
        add_section("4. Academic Orientation & Future Aspirations", student.report.academic_future_outlook)
        add_section("5. Values & Character Development Alignment", student.report.institutional_value_alignment)

        # Observations Rubric Table
        h_obs = doc.add_paragraph()
        h_obs.paragraph_format.space_before = Pt(10)
        h_obs.paragraph_format.space_after = Pt(3)
        r_obs = h_obs.add_run("6. Clinical Observation Scorecard")
        r_obs.font.name = "Georgia"
        r_obs.font.size = Pt(11.5)
        r_obs.font.bold = True
        r_obs.font.color.rgb = EMERALD

        obs_table = doc.add_table(rows=6, cols=2)
        obs_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        obs_pairs = [
            ("Emotional Expression", student.observations.emotional_expression, "Confidence Level", student.observations.confidence_level),
            ("Anxiety / Restlessness", student.observations.anxiety_restlessness, "Clarity vs Confusion", student.observations.clarity_vs_confusion),
            ("Family Pressure Themes", student.observations.family_pressure_themes, "Fear of Failure/Judgement", student.observations.fear_of_failure_judgement),
            ("Motivational Level", student.observations.motivational_level, "Identity Struggles", student.observations.identity_struggles),
            ("Peer Influence", student.observations.peer_influence, "Emotional Maturity", student.observations.emotional_maturity),
            ("Triage Recommendation", student.report.triage_status, "Follow-up Needed", "Yes" if student.report.follow_up_recommended else "No")
        ]

        for i, (k1, val1, k2, val2) in enumerate(obs_pairs):
            r = obs_table.rows[i]
            c1, c2 = r.cells[0], r.cells[1]
            set_cell_background(c1, "F3F4F6" if i % 2 == 0 else "FFFFFF")
            set_cell_background(c2, "F3F4F6" if i % 2 == 0 else "FFFFFF")
            set_cell_margins(c1, top=40, bottom=40, left=80, right=80)
            set_cell_margins(c2, top=40, bottom=40, left=80, right=80)

            p1 = c1.paragraphs[0]
            p1.add_run(f"• {k1}: ").bold = True
            p1.add_run(str(val1))
            p1.runs[0].font.size = Pt(9)
            p1.runs[1].font.size = Pt(9)

            p2 = c2.paragraphs[0]
            p2.add_run(f"• {k2}: ").bold = True
            p2.add_run(str(val2))
            p2.runs[0].font.size = Pt(9)
            p2.runs[1].font.size = Pt(9)

        add_section("7. Counsellor's Professional Remarks", student.report.counsellor_remarks)
        add_section("8. Prescribed Exercises & Recommendations", student.report.prescribed_exercises)
        
        if student.report.follow_up_notes:
            add_section("9. Pastoral / Follow-up Action Plan", student.report.follow_up_notes)

    # Signatures
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(24)
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sig = p_sig.add_run("_______________________________\nCounsellor Signature & Date\nAdolescent Guidance Department")
    r_sig.font.size = Pt(9.5)
    r_sig.font.color.rgb = DARK_GRAY

    target_io = io.BytesIO()
    doc.save(target_io)
    target_io.seek(0)
    return target_io


def generate_cumulative_docx(cumulative: CumulativeReport) -> io.BytesIO:
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = title_p.add_run("STUDENT COUNSELLING & GUIDANCE DEPARTMENT\n")
    run_inst.font.name = "Georgia"
    run_inst.font.size = Pt(17)
    run_inst.font.bold = True
    run_inst.font.color.rgb = EMERALD

    run_sub = title_p.add_run("ADOLESCENT COHORT EVALUATION & GUIDANCE DOSSIER\n")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = GOLD

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(f"Report Generated: {cumulative.generated_at} | Total Candidates Evaluated: {cumulative.total_evaluated}\n")
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = DARK_GRAY

    p_div = doc.add_paragraph()
    p_div_run = p_div.add_run("―" * 58)
    p_div_run.font.color.rgb = GOLD
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Stats Summary Table
    t_stats = doc.add_table(rows=1, cols=3)
    t_stats.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cells = t_stats.rows[0].cells
    labels = [
        (f"🟢 Well-Adapted\n{cumulative.triage_counts.get('Well-Adapted', 0)} Students", GREEN_BG, RGBColor(21, 128, 61)),
        (f"🟡 Mild Observation\n{cumulative.triage_counts.get('Mild Observation', 0)} Students", AMBER_BG, RGBColor(180, 130, 20)),
        (f"🔴 Requires Follow-up\n{cumulative.triage_counts.get('Requires Follow-up', 0)} Students", RED_BG, RGBColor(185, 28, 28))
    ]

    for idx, (txt, bg, col) in enumerate(labels):
        cell = cells[idx]
        set_cell_background(cell, bg)
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt)
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = col

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Executive Overview
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    r_h1 = h1.add_run("1. Executive Pastoral Overview")
    r_h1.font.bold = True
    r_h1.font.size = Pt(12)
    r_h1.font.color.rgb = EMERALD

    p_exec = doc.add_paragraph(cumulative.executive_overview)
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.paragraph_format.space_after = Pt(8)

    # Thematic Analysis
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    r_h2 = h2.add_run("2. Prevalent Cohort Psychological & Developmental Themes")
    r_h2.font.bold = True
    r_h2.font.size = Pt(12)
    r_h2.font.color.rgb = EMERALD

    for item in cumulative.key_cohort_themes:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_after = Pt(4)
        r_th = p_t.add_run(f"• {item['theme']}: ")
        r_th.bold = True
        r_th.font.size = Pt(10.5)
        r_obs = p_t.add_run(item['observation'])
        r_obs.font.size = Pt(10.5)

    # Spiritual Integration
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    r_h3 = h3.add_run("3. Spiritual Assimilation & Practice of Moula's Guidance")
    r_h3.font.bold = True
    r_h3.font.size = Pt(12)
    r_h3.font.color.rgb = EMERALD

    p_sp = doc.add_paragraph(cumulative.spiritual_value_integration_summary)
    p_sp.paragraph_format.line_spacing = 1.15
    p_sp.paragraph_format.space_after = Pt(8)

    # High-Risk Action Roster
    if cumulative.action_roster_high_risk:
        h4 = doc.add_paragraph()
        h4.paragraph_format.space_before = Pt(12)
        r_h4 = h4.add_run("4. Action Roster: Candidates Requiring Continued Guidance")
        r_h4.font.bold = True
        r_h4.font.size = Pt(12)
        r_h4.font.color.rgb = EMERALD

        t_roster = doc.add_table(rows=len(cumulative.action_roster_high_risk) + 1, cols=4)
        t_roster.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["Candidate Name", "Triage Level", "Key Area of Concern", "Recommended Next Step"]
        for c_idx, h_text in enumerate(headers):
            cell = t_roster.rows[0].cells[c_idx]
            set_cell_background(cell, "E5E7EB")
            p = cell.paragraphs[0]
            r = p.add_run(h_text)
            r.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = DARK_GRAY

        for r_idx, c_info in enumerate(cumulative.action_roster_high_risk, start=1):
            row = t_roster.rows[r_idx]
            bg = RED_BG if "Requires Follow-up" in c_info.get("triage", "") else AMBER_BG
            
            data_row = [
                c_info.get("name", ""),
                c_info.get("triage", ""),
                c_info.get("key_concern", ""),
                c_info.get("action_recommended", "")
            ]
            for col_idx, val in enumerate(data_row):
                cell = row.cells[col_idx]
                set_cell_background(cell, bg if col_idx == 1 else "FAFAFA")
                set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
                p = cell.paragraphs[0]
                r = p.add_run(str(val))
                r.font.size = Pt(9)

    # Management Recommendations
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(14)
    r_h5 = h5.add_run("5. Recommendations for Program Management & Mentors")
    r_h5.font.bold = True
    r_h5.font.size = Pt(12)
    r_h5.font.color.rgb = EMERALD

    for rec in cumulative.management_recommendations:
        p_r = doc.add_paragraph()
        p_r.paragraph_format.space_after = Pt(4)
        r = p_r.add_run(rec)
        r.font.size = Pt(10.5)

    # Signatures
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(30)
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sig = p_sig.add_run("_______________________________\nLead Counsellor & Program Director\nStudent Guidance Department")
    r_sig.font.size = Pt(9.5)
    r_sig.font.color.rgb = DARK_GRAY

    target_io = io.BytesIO()
    doc.save(target_io)
    target_io.seek(0)
    return target_io

